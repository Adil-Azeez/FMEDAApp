"""Regression coverage for descriptions and the Project Information fields."""
import json
from unittest.mock import patch

import openpyxl
import pytest
from PyQt6.QtWidgets import QLabel, QMessageBox, QTextEdit
from PyQt6.QtGui import QTextDocument

from fmeda_tool.db.database import get_db_connection, ensure_database_ready
from fmeda_tool.models import Project, SafetyContext, Unit, Component, FailureModeAssignment
from fmeda_tool.services.component_library_service import ComponentLibraryService as Library
from fmeda_tool.services.project_service import ProjectService
from fmeda_tool.services.export_service import ExportService
from fmeda_tool.services.calculation_service import CalculationService
from fmeda_tool.ui.create_project_view import CreateProjectView
from fmeda_tool.ui.components_db_view import ComponentsDBView
from fmeda_tool.ui.dialogs.custom_component_dialog import CustomComponentDialog
from fmeda_tool.ui.dialogs.component_mapping_dialog import ComponentMappingDialog
from fmeda_tool.ui.unit_editor_view import ProjectOverviewTab


DESCRIPTION = "Temperature sensor\nUsed in channel A"
ARCHITECTURE = "Two independent channels <A & B>\nShared output monitoring"
SCHEMATIC = "  SCH-02 / Rev. B + ä  "


def create_component(description=DESCRIPTION):
    return Library.create_custom_component(
        component_type="Sensor", fits=12.5, failure_modes={"Open": 40, "Short": 60},
        description=description,
    )


def project():
    return Project(
        id="fields", name="Fields", description="Test", project_number="F-01",
        product_version="product", hardware_version="hardware", software_version="software",
        schematic_version=SCHEMATIC, selected_profile="Profile 4", environmental_profile="Profile 4",
        safety_context=SafetyContext(safety_architecture=ARCHITECTURE,
                                     operating_mode="Continuous mode", external_sensor_included=True),
    )


@pytest.mark.parametrize("description", ["", " \n\t ", None])
def test_service_rejects_blank_description(fresh_isolated_db, description):
    ok, message, snapshot = create_component(description)
    assert not ok and "Description is required" in message and snapshot is None
    assert Library.search_custom_components() == []


def test_create_edit_and_sqlite_persistence(fresh_isolated_db, qapp):
    dialog = CustomComponentDialog()
    assert isinstance(dialog.description_input, QTextEdit)
    dialog.type_input.setText("Temperature Sensor")
    dialog.description_input.setPlainText(" \n" + DESCRIPTION + " \n")
    with patch.object(QMessageBox, "information"):
        dialog._on_save()
    snapshot = dialog.saved_snapshot
    assert snapshot["description"] == DESCRIPTION
    cid = snapshot["library_component_id"]
    with get_db_connection() as conn:
        assert conn.execute("SELECT description FROM custom_components WHERE id = ?", (cid,)).fetchone()[0] == DESCRIPTION
    assert Library.search_custom_components("channel A")[0]["id"] == cid
    edit = CustomComponentDialog(component_id=cid)
    assert edit.description_input.toPlainText() == DESCRIPTION
    edit.description_input.setPlainText("  Revised\ndescription  ")
    with patch.object(QMessageBox, "information"):
        edit._on_save()
    assert Library.get_custom_component_snapshot(cid)["description"] == "Revised\ndescription"
    assert Library.get_custom_component_snapshot(cid)["failure_modes"] == snapshot["failure_modes"]


@pytest.mark.parametrize("is_edit", [False, True])
@pytest.mark.parametrize("description", ["", " \n\t "])
def test_dialog_rejects_blank_description(fresh_isolated_db, qapp, is_edit, description):
    cid = create_component()[2]["library_component_id"] if is_edit else None
    dialog = CustomComponentDialog(component_id=cid)
    dialog.type_input.setText("Sensor")
    dialog.description_input.setPlainText(description)
    with patch.object(QMessageBox, "warning") as warning:
        dialog._on_save()
    assert "Description is required" in warning.call_args.args[2]
    assert dialog.saved_snapshot is None
    if is_edit:
        assert Library.get_custom_component_snapshot(cid)["description"] == DESCRIPTION


def test_old_database_migrates_in_place_and_requires_description_on_edit(fresh_isolated_db, qapp):
    cid = create_component()[2]["library_component_id"]
    with get_db_connection() as conn:
        source_rows = conn.execute("SELECT * FROM components").fetchall()
        legacy_rows = conn.execute("SELECT * FROM legacy_components").fetchall()
        custom_row = dict(conn.execute("SELECT * FROM custom_components WHERE id = ?", (cid,)).fetchone())
        conn.execute("ALTER TABLE custom_components DROP COLUMN description")
    assert ensure_database_ready()[0]
    assert ensure_database_ready()[0]  # Migration is idempotent.
    with get_db_connection() as conn:
        assert conn.execute("SELECT * FROM components").fetchall() == source_rows
        assert conn.execute("SELECT * FROM legacy_components").fetchall() == legacy_rows
        migrated = dict(conn.execute("SELECT * FROM custom_components WHERE id = ?", (cid,)).fetchone())
        assert migrated.pop("description") == ""
        custom_row.pop("description")
        assert migrated == custom_row
    assert Library.get_custom_component_snapshot(cid)["description"] == ""
    dialog = CustomComponentDialog(component_id=cid)
    with patch.object(QMessageBox, "warning") as warning:
        dialog._on_save()
    assert "Description is required" in warning.call_args.args[2]
    ok, message, _ = Library.update_custom_component(cid, "Sensor", 12.5, {"Open": 40, "Short": 60})
    assert not ok and "Description is required" in message
    dialog.description_input.setPlainText(DESCRIPTION)
    with patch.object(QMessageBox, "information"):
        dialog._on_save()
    assert Library.get_custom_component_snapshot(cid)["description"] == DESCRIPTION


@pytest.mark.parametrize("source", ["exida", "legacy"])
def test_copy_description_without_modifying_source(fresh_isolated_db, qapp, source):
    search = Library.search_exida_components if source == "exida" else Library.search_legacy_components
    snapshot_getter = Library.get_exida_component_snapshot if source == "exida" else Library.get_legacy_component_snapshot
    row = next(r for r in search() if (r.get("fit") or r.get("fits") or 0) > 0 and len(r["failure_modes"]) >= 2)
    # Optional descriptions in source catalogs must be preferred over descriptive labels.
    table = "components" if source == "exida" else "legacy_components"
    with get_db_connection() as conn:
        conn.execute(f"ALTER TABLE {table} ADD COLUMN description TEXT")
        conn.execute(f"UPDATE {table} SET description = ? WHERE id = ?", (DESCRIPTION, row["id"]))
    before = snapshot_getter(row["id"])
    dialog = CustomComponentDialog()
    dialog._apply_copied_snapshot(before)
    assert dialog.description_input.toPlainText() == DESCRIPTION
    dialog.display_name_input.clear()
    dialog.description_input.setPlainText(DESCRIPTION + "\nEdited copy")
    with patch.object(QMessageBox, "information"):
        dialog._on_save()
    assert dialog.saved_snapshot["description"] == DESCRIPTION + "\nEdited copy"
    assert snapshot_getter(row["id"]) == before
    assert dialog.saved_snapshot["copied_from_component_id"] == row["id"]


def test_description_visible_in_library_and_mapping(fresh_isolated_db, qapp):
    ok, message, snapshot = create_component()
    assert ok, message
    view = ComponentsDBView()
    view.custom_search_input.setText("channel A")
    assert view.custom_table.rowCount() == 1
    assert view.custom_table.item(0, 7).text() == DESCRIPTION
    mapping = ComponentMappingDialog(Unit(id="u", name="Group", description="Test"))
    mapping.custom_radio.setChecked(True)
    assert mapping.lib_table.item(0, 4).text() == DESCRIPTION
    mapping.lib_table.selectRow(0)
    mapping._on_lib_table_selection_changed()
    assert mapping.cfg_description.text() == DESCRIPTION
    assert mapping.selected_template.snapshot["description"] == DESCRIPTION
    view.close()
    mapping.close()


def test_project_information_fields_and_legacy_values(qapp):
    original = project()
    original.notes = "Retain notes"
    original.change_history = [{"action": "Existing history"}]
    view = CreateProjectView()
    view.load_project(original.model_dump(mode="json"))
    labels = [label.text() for label in view.findChildren(QLabel)]
    assert not any("Operating Mode" in label or "External Sensor Included" in label for label in labels)
    assert isinstance(view.architecture_input, QTextEdit)
    assert view.architecture_input.toPlainText() == ARCHITECTURE
    assert view.schematic_version_input.text() == SCHEMATIC
    product_row, _ = view.form.getWidgetPosition(view.product_version_input)
    schematic_row, _ = view.form.getWidgetPosition(view.schematic_version_input)
    assert schematic_row == product_row + 1
    view.architecture_input.setPlainText(ARCHITECTURE + "\nNew line")
    view.schematic_version_input.setText("  Exact / schematic  ")
    view._on_next()
    saved = view.project
    assert saved.schematic_version == "  Exact / schematic  "
    assert saved.safety_context.safety_architecture == ARCHITECTURE + "\nNew line"
    assert saved.safety_context.operating_mode == "Continuous mode"
    assert saved.safety_context.external_sensor_included is True
    for field in ("selected_profile", "environmental_profile", "product_version", "hardware_version", "software_version", "notes", "change_history"):
        assert getattr(saved, field) == getattr(original, field)
    view.reset_form()
    assert view.schematic_version_input.text() == ""
    view.name_input.setText("New project")
    view.number_input.setText("N-01")
    view.description_input.setPlainText("New description")
    view.architecture_input.setPlainText(ARCHITECTURE)
    view.schematic_version_input.setText(SCHEMATIC)
    view._on_next()
    assert view.project.safety_context.safety_architecture == ARCHITECTURE
    assert view.project.schematic_version == SCHEMATIC


@pytest.mark.parametrize("schema_version", [1, 2])
@pytest.mark.parametrize("architecture", ["1oo2", "2oo2", ARCHITECTURE])
def test_old_projects_migrate_and_roundtrip(tmp_path, schema_version, architecture):
    data = project().model_dump(mode="json")
    data.pop("schematic_version")
    data["schema_version"] = schema_version
    data["safety_context"]["safety_architecture"] = architecture
    path = tmp_path / "old.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    loaded, migrated, _ = ProjectService.load_and_migrate_project(str(path))
    assert migrated == (schema_version == 1)
    assert loaded.schematic_version == ""
    assert loaded.safety_context.safety_architecture == architecture
    ProjectService.save_project_atomically(loaded, str(path))
    expected = dict(data, schema_version=2, schematic_version="")
    assert json.loads(path.read_text(encoding="utf-8")) == expected
    loaded.schematic_version = SCHEMATIC
    loaded.safety_context.safety_architecture = ARCHITECTURE
    ProjectService.save_project_atomically(loaded, str(path))
    reopened, _, _ = ProjectService.load_and_migrate_project(str(path))
    assert reopened.model_dump() == loaded.model_dump()
    assert json.loads(path.with_suffix(".json.bak").read_text(encoding="utf-8")) == expected


def test_overview_and_exports_contain_values(tmp_path, qapp):
    p = project()
    overview = ProjectOverviewTab(None)
    overview.refresh(p)
    texts = [label.text() for label in overview.findChildren(QLabel)]
    assert ARCHITECTURE in texts and SCHEMATIC in texts
    xlsx = tmp_path / "overview.xlsx"
    assert ExportService.export_to_excel(p, str(xlsx))
    wb = openpyxl.load_workbook(xlsx)
    ws = wb["Overview"]
    cells = {cell.value: cell for row in ws for cell in row if isinstance(cell.value, str)}
    assert ws.cell(cells["Schematic Version:"].row, 2).value == SCHEMATIC
    assert ws.cell(cells["Safety Architecture:"].row, 5).value == ARCHITECTURE
    assert ws.cell(cells["Safety Architecture:"].row, 5).alignment.wrap_text
    wb.close()
    # Inspect the exact document sent to the real Qt PDF printer, including line breaks.
    captured = []
    original_set_html = QTextDocument.setHtml
    def capture(document, html):
        original_set_html(document, html)
        captured.append(document.toPlainText())
    pdf = tmp_path / "overview.pdf"
    with patch.object(QTextDocument, "setHtml", capture):
        assert ExportService.export_to_pdf(p, str(pdf))
    assert pdf.read_bytes().startswith(b"%PDF")
    assert "Schematic Version:" in captured[0]
    assert SCHEMATIC.strip() in captured[0]
    assert ARCHITECTURE in captured[0]


@pytest.mark.parametrize("architecture", ["1oo1", "1oo2", "2oo2"])
def test_calculations_identical_after_project_roundtrip(tmp_path, architecture):
    p = project()
    p.safety_context.safety_architecture = architecture
    component = Component(id="c", position="U1", name="Sensor", type="Sensor", failure_rate=10,
                          failure_modes={"Open": 100}, failure_mode_assignments=[
                              FailureModeAssignment(failure_mode_name="Open", failure_rate_percentage=100,
                                                    classification="dangerous_failure", detection_percentage=80)
                          ])
    p.units = [Unit(id="u", name="Group", description="Test", components=[component])]
    before = CalculationService.calculate_scope(p.units, p)
    p.schematic_version = "Changed metadata"
    path = tmp_path / "calc.json"
    ProjectService.save_project_atomically(p, str(path))
    loaded, _, _ = ProjectService.load_and_migrate_project(str(path))
    assert CalculationService.calculate_scope(loaded.units, loaded) == before
