"""Classification review, isolated persistence, and independent result regression tests."""
import hashlib
import inspect
import json
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pytest
from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import QHeaderView

from fmeda_tool.models import Project, Unit, Component, FailureModeAssignment
from fmeda_tool.models.failure_mode_assignment import CLASSIFICATIONS, LEGACY_CLASSIFICATION
from fmeda_tool.services.calculation_service import CalculationService as Calc
from fmeda_tool.services.project_service import ProjectService
from fmeda_tool.services.validation_service import ValidationService
from fmeda_tool.services.export_service import ExportService
from fmeda_tool.ui.models.fmeda_table_model import FmedaTableModel, COLUMN_HEADERS
from fmeda_tool.ui.main_window import MainWindow
from fmeda_tool.ui.verification_view import VerificationView
from fmeda_tool.ui.dialogs.component_instance_dialog import ComponentInstanceDialog


def make_project():
    units = []
    for index, (classification, fit) in enumerate((('dangerous_failure', 100), ('safe_failure', 40), ('no_part_failure', 70))):
        assignment = FailureModeAssignment(failure_mode_name='Open', failure_rate_percentage=100,
            classification=classification, diagnostic_function='Yes', detection_percentage=80,
            proof_test_a=50, proof_test_b=25, proof_test_c=10)
        component = Component(id=f'c{index}', name='Resistor', type='Resistor', position=f'R{index}',
            failure_rate=fit, failure_modes={'Open': 100}, failure_mode_assignments=[assignment])
        units.append(Unit(id=f'u{index}', name=f'Group {index}', description='Independent assessment',
                          components=[component], included_in_safety_function=index != 1))
    return Project(id='p', name='Classification', description='Isolated tests', units=units)


@pytest.mark.parametrize('classification, bucket', [
    ('safe_failure', 'lambda_safe'), ('dangerous_failure', 'lambda_dangerous'),
    ('no_part_failure', 'lambda_no_part'), ('no_effect_failure', 'lambda_no_effect'),
    ('diagnostic_function_failure', 'lambda_diag_failure'), ('not_evaluated', 'lambda_unclassified')])
def test_exclusive_allocation(classification, bucket):
    result = Calc.calculate_row_detailed(100, classification, 37, 80)
    buckets = ('lambda_safe', 'lambda_dangerous', 'lambda_no_part', 'lambda_no_effect', 'lambda_diag_failure', 'lambda_unclassified')
    assert result[bucket] == 100
    assert sum(result[k] for k in buckets) == 100
    if classification == 'safe_failure':
        assert result['lambda_sd'] == 80 and result['lambda_su'] == pytest.approx(20)
    elif classification == 'dangerous_failure':
        assert result['lambda_dd'] == 80 and result['lambda_du'] == pytest.approx(20)
    else:
        assert all(result[k] == 0 for k in ('lambda_sd', 'lambda_su', 'lambda_dd', 'lambda_du'))


def test_selection_workflows_and_independent_diagnostic_flag(qtbot):
    p = make_project()
    a = p.units[0].components[0].failure_mode_assignments[0]
    a.set_classification('not_evaluated')
    model = FmedaTableModel(p.units[0], p)
    assert len(model._get_dropdown_options(11)) == 6
    assert 'No Part / No Effect' not in COLUMN_HEADERS
    assert not model.setData(model.index(0, 11), 'safe_failure')
    model.set_edit_mode(True)
    for classification in CLASSIFICATIONS:
        assert model.setData(model.index(0, 11), classification)
        assert a.diagnostic_function == 'Yes'
        assert Calc.calculate_diagnostic_lambda(p.units[0].components[0]) == 100
        assert model.search_rows(CLASSIFICATIONS[classification])[0] == [0]
    dialog = ComponentInstanceDialog(p.units[0].components[0], p)
    qtbot.addWidget(dialog)
    assert dict((value, label) for label, value in dialog._choices('classification')) == CLASSIFICATIONS
    assert 'dont_care' not in dialog.ENGINEERING_FIELDS
    assert FailureModeAssignment(failure_mode_name='New', failure_rate_percentage=100).classification == 'not_evaluated'


def legacy_file(tmp_path):
    p = make_project()
    payload = p.model_dump(mode='json')
    for unit in payload['units']:
        a = unit['components'][0]['failure_mode_assignments'][0]
        a.pop('classification_version')
        a['classification'] = 'dangerous_failure'
        a['dangerous_failure_percentage'] = 60
    payload['units'][0]['components'][0]['failure_mode_assignments'][0]['dont_care'] = True
    path = tmp_path / 'legacy.json'
    path.write_text(json.dumps(payload), encoding='utf-8')
    return path


def test_migration_read_only_review_undo_and_roundtrip(tmp_path, qtbot):
    path = legacy_file(tmp_path)
    original = path.read_bytes()
    project, migrated, message = ProjectService.load_and_migrate_project(str(path))
    assert migrated and '1 rows' in message
    a = project.units[0].components[0].failure_mode_assignments[0]
    assert a.classification == LEGACY_CLASSIFICATION and a.dont_care
    assert project.units[1].components[0].failure_mode_assignments[0].classification == 'dangerous_failure'
    metrics = Calc.calculate_component(project.units[0].components[0])
    assert metrics['gesamtgerat']['lambda_dd'] == 48
    assert metrics['gesamtgerat']['lambda_safe'] == 40
    assert metrics['sicherheitskanal']['lambda'] == 0
    other = Calc.calculate_component(project.units[1].components[0])
    assert other['gesamtgerat']['lambda_dd'] == pytest.approx(19.2)
    findings = ValidationService.validate_project(project)
    review = next(f for f in findings if 'Classification Review Required' in f['message'])
    assert (review['unit_id'], review['row_index']) == ('u0', 0)
    window = MainWindow()
    qtbot.addWidget(window)
    window.current_project = project
    window.unit_editor_view.load_project(project)
    window.unit_editor_view.unit_tabs.setCurrentIndex(1)
    assert not window.has_unsaved_changes and not window.undo_stack
    assert path.read_bytes() == original
    tab = window.unit_editor_view.unit_tabs.widget(1)
    tab.enable_editing()
    assert not tab.model.setData(tab.model.index(0, 11), 'safe_failure')
    with patch.object(Calc, 'calculate_project', wraps=Calc.calculate_project) as calculate:
        assert tab.model.setData(tab.model.index(0, 11), 'no_effect_failure')
        assert calculate.call_count == 1
    tab.confirm_changes()
    assert len(window.undo_stack) == 1
    assert a.dont_care and a.classification_version == 2
    assert Calc.calculate_group_result(project.units[0], project)['lambda_no_effect'] == 100
    assert not any('Classification Review Required' in f['message'] for f in ValidationService.validate_project(project))
    destination = tmp_path / 'reviewed.json'
    ProjectService.save_project_atomically(project, str(destination))
    reopened, _, _ = ProjectService.load_and_migrate_project(str(destination))
    assert reopened.units[0].components[0].failure_mode_assignments[0].classification == 'no_effect_failure'
    assert path.read_bytes() == original
    window._on_undo()
    assert window.current_project.units[0].components[0].failure_mode_assignments[0].classification == LEGACY_CLASSIFICATION


def test_independent_matrix_excel_and_undefined_values(tmp_path, qtbot):
    p = make_project()
    view = VerificationView()
    qtbot.addWidget(view)
    before = p.model_dump_json()
    view.load_project(p)
    assert p.model_dump_json() == before
    assert view.secondary_table.columnCount() == 3
    assert view.card_frame_sil.isHidden()
    assert [r['lambda_total'] for r in view.group_results] == [100, 40, 70]
    assert [r['sff'] for r in view.group_results] == [80, 100, None]
    assert view.group_results[2]['achieved_sil'] == 'N/A'
    assert view.group_results[1]['pfd_avg'] == 0
    assert view.group_results[0]['proof_test_a_contribution'] == pytest.approx(10)
    table = view.secondary_table
    assert table.horizontalHeader().sectionResizeMode(0) == QHeaderView.ResizeMode.Interactive
    table.setColumnWidth(0, 240)
    view._update_scope_results()
    assert table.columnWidth(0) == 240
    view.resize(650, 750)
    view.show()
    qtbot.waitUntil(lambda: table.horizontalScrollBar().maximum() > 0)
    book_path = tmp_path / 'independent.xlsx'
    assert ExportService.export_to_excel(p, str(book_path))
    book = openpyxl.load_workbook(book_path)
    sheet = book['Overview']
    title = next(c for row in sheet for c in row if c.value == 'Calculated Safety Metrics by Functional Group')
    header = title.row + 2
    for col, unit in enumerate(p.units):
        assert sheet.cell(header + col + 1, 1).value == unit.name
        for row, (key, label, style) in enumerate(Calc.GROUP_METRICS):
            raw = table.item(row, col).data(Qt.ItemDataRole.UserRole)
            expected = 'N/A' if raw is None else raw / 100 if style == 'pct' else raw
            actual = sheet.cell(header + col + 1, row + 2).value
            assert actual == pytest.approx(expected) if isinstance(expected, float) else actual == expected
    assert sheet.max_row == header + 3
    assert not any(c.value in ('Selected Scope Total', 'Calculated Safety Metrics Summary') for row in sheet for c in row)
    headers = [c.value for c in book['Group 2'][1]]
    assert 'No Part / No Effect' not in headers
    assert book['Group 2'].cell(2, headers.index('Failure Classification') + 1).value == 'No Part Failure'
    assert book['Group 2'].cell(2, headers.index('lambda_no_part (FIT)') + 1).value == 70
    book.close()
    view.selected_unit_ids = {'u1'}
    view._update_scope_results()
    assert table.columnCount() == 1 and not view.card_frame_sff.isHidden()
    assert view.card_val_sff.text() == '100.00%'


def test_legacy_review_export_and_unclassified_warning(tmp_path):
    p, _, _ = ProjectService.load_and_migrate_project(str(legacy_file(tmp_path)))
    p.units[1].components[0].failure_mode_assignments[0].set_classification('not_evaluated')
    findings = ValidationService.validate_project(p)
    assert any(f['unit_id'] == 'u1' and f['row_index'] == 0 and f['message'] == 'Failure Classification not evaluated' for f in findings)
    path = tmp_path / 'review.xlsx'
    assert ExportService.export_to_excel(p, str(path))
    book = openpyxl.load_workbook(path)
    assert 'Classification Review Required' in book['Group 0'].cell(2, 12).value
    book.close()


def test_pdf_implementation_is_unchanged():
    source = Path(inspect.getsourcefile(ExportService)).read_text(encoding="utf-8")
    pdf = source[source.index("    def export_to_pdf("):]
    assert hashlib.sha256(pdf.encode()).hexdigest() == '0195b9909f1f796706063fc28c91165c54a9830954540b2938f1b9708d8db577'

@pytest.mark.parametrize('classification', ['no_part_failure', 'no_effect_failure', 'diagnostic_function_failure'])
def test_excluded_classes_do_not_change_safety_metrics(classification):
    p = make_project()
    unit = p.units[0]
    before = Calc.calculate_group_result(unit, p)
    excluded = unit.components[0].model_copy(deep=True)
    excluded.id = 'excluded'
    excluded.failure_rate = 10000
    excluded.failure_mode_assignments[0].set_classification(classification)
    excluded.failure_mode_assignments[0].diagnostic_function = 'No'
    unit.components.append(excluded)
    after = Calc.calculate_group_result(unit, p)
    for key in ('sff', 'dc', 'mttfd', 'pfd_avg', 'pfhd', 'achieved_sil', 'lambda_diag_gesamt'):
        assert after[key] == before[key]


def test_bulk_classification_and_component_dialog_review(qtbot, tmp_path):
    from fmeda_tool.ui.unit_editor_view import BulkActionSummaryDialog
    from PyQt6.QtWidgets import QDialog
    p = make_project()
    window = MainWindow()
    qtbot.addWidget(window)
    window.current_project = p
    window.unit_editor_view.load_project(p)
    window.unit_editor_view.unit_tabs.setCurrentIndex(1)
    tab = window.unit_editor_view.unit_tabs.widget(1)
    tab.enable_editing()
    tab.table.selected_component_ids = {'c0'}
    with patch.object(BulkActionSummaryDialog, 'exec', return_value=QDialog.DialogCode.Accepted):
        for classification in CLASSIFICATIONS:
            assert tab.bulk_set_classification(classification)
            assert p.units[0].components[0].failure_mode_assignments[0].classification == classification
    legacy, _, _ = ProjectService.load_and_migrate_project(str(legacy_file(tmp_path)))
    comp = legacy.units[0].components[0]
    dialog = ComponentInstanceDialog(comp, legacy)
    qtbot.addWidget(dialog)
    combo = dialog._engineering_widget(0, 'classification')
    assert 'Classification Review Required' in combo.currentText()
    combo.setCurrentIndex(combo.findData('no_part_failure'))
    dialog._on_save()
    assert dialog.result() == QDialog.DialogCode.Accepted
    assert comp.failure_mode_assignments[0].classification == 'no_part_failure'
    assert comp.failure_mode_assignments[0].dont_care


def test_new_component_modes_never_guess_classification(qtbot):
    from fmeda_tool.ui.dialogs.component_selection_dialog import ComponentSelectionDialog
    dialog = ComponentSelectionDialog()
    qtbot.addWidget(dialog)
    dialog._load_failure_modes({'Open': 25, 'Short': 25, 'Drift': 25, 'Loss': 25})
    assert all(combo.currentText() == 'Not Evaluated' for combo in dialog.comboboxes.values())


def test_page3_never_requests_combined_calculation(qtbot):
    view = VerificationView()
    qtbot.addWidget(view)
    with patch.object(Calc, 'calculate_scope', wraps=Calc.calculate_scope) as scope:
        view.load_project(make_project())
        view._clear_all_groups()
        view._select_all_groups()
    assert all(len(call.args[0]) == 1 for call in scope.call_args_list)
