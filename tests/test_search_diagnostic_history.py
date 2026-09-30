"""Isolated regressions for FMEDA search, diagnostic FIT and successful-save audits."""
import json
from contextlib import ExitStack
from copy import deepcopy
from datetime import datetime
from unittest.mock import patch

import openpyxl
import pytest
from PyQt6.QtCore import Qt
from PyQt6.QtGui import QTextDocument
from PyQt6.QtWidgets import QAbstractItemView, QComboBox, QInputDialog, QMessageBox, QStyleOptionViewItem

from fmeda_tool.models import Project, Unit, Component, FailureModeAssignment
from fmeda_tool.services.calculation_service import CalculationService
from fmeda_tool.services.validation_service import ValidationService
from fmeda_tool.services.project_service import ProjectService
from fmeda_tool.services.change_history_service import ChangeHistoryService
from fmeda_tool.services.export_service import ExportService
from fmeda_tool.ui.models.fmeda_table_model import FmedaTableModel
from fmeda_tool.ui.unit_editor_view import UnitEditorView, ChangeHistoryDialog
from fmeda_tool.ui.verification_view import VerificationView
from fmeda_tool.ui.main_window import MainWindow


def make_project():
    components = []
    for i, position in enumerate(("R1", "U2")):
        components.append(Component(
            id=f"c{i}", position=position, name="Precision sensor", type="Amplifier",
            failure_rate=20.0, value="Signal path", internal_pn=f"PN-{i}",
            failure_modes={"Open": 40.0, "Short": 60.0},
            failure_mode_assignments=[FailureModeAssignment(
                failure_mode_name=name, failure_rate_percentage=pct, classification="safe_failure",
                dangerous_failure_percentage=0.0, detection_percentage=0.0,
                dc_test_ref=f"HiddenRef-{i}", notes="Row justification", review_status="approved",
            ) for name, pct in {"Open": 40.0, "Short": 60.0}.items()],
        ))
    return Project(id="audit", name="Audit Project", description="Description", created_by="Engineer",
                   units=[Unit(id="u1", name="Sensor group", description="First", components=components),
                          Unit(id="u2", name="Other group", description="Second", components=[
                              Component(id="c3", position="C3", name="Cap", type="Capacitor", failure_rate=10,
                                        failure_modes={"Drift": 100}, failure_mode_assignments=[
                                            FailureModeAssignment(failure_mode_name="Drift", failure_rate_percentage=100,
                                                                  diagnostic_function="Yes", dont_care=True)
                                        ])])])


@pytest.mark.parametrize("term, expected", [("r1", [0, 1]), ("PRECIS", [0, 1, 3, 4]),
    ("Ampl", [0, 1, 3, 4]), ("signal", [0, 1, 3, 4]), ("pn-0", [0, 1]),
    ("oPeN", [0, 3]), ("HiddenRef-1", [3, 4]), ("approved", [0, 1, 3, 4]),
    ("20.0000", [0, 1, 3, 4]), ("justif", [0, 1, 3, 4]), ("safe failure", [0, 1, 3, 4])])
def test_model_search_all_columns(qapp, term, expected):
    p = make_project()
    model = FmedaTableModel(p.units[0], p)
    original = p.model_dump()
    with patch.object(CalculationService, "calculate_row_detailed", side_effect=AssertionError("Search recalculated")), \
         patch.object(ValidationService, "validate_row", side_effect=AssertionError("Search validated")):
        matches, visible_rows = model.search_rows(term)
    assert matches == expected
    assert visible_rows == set(expected)
    assert len(matches) == len(set(matches))
    assert p.model_dump() == original


@pytest.mark.parametrize("edit_mode", [False, True])
def test_search_controls_row_filter_and_lazy_tabs(qapp, edit_mode):
    p = make_project()
    editor = UnitEditorView()
    editor.load_project(p)
    editor.unit_tabs.setCurrentIndex(1)
    tab, other = editor.unit_tabs.widget(1), editor.unit_tabs.widget(2)
    if edit_mode:
        tab.enable_editing()
    p.units[0].components[0].failure_mode_assignments[0].notes = "unique needle"
    data = deepcopy(p.model_dump())
    emitted = []
    editor.project_changed.connect(lambda: emitted.append(True))
    with patch.object(CalculationService, "calculate_project", side_effect=AssertionError("Search recalculated")), \
         patch.object(ValidationService, "validate_row", side_effect=AssertionError("Search validated")), \
         patch.object(other.model, "reload_data", side_effect=AssertionError("Unrelated tab rebuilt")):
        search = tab.search_bar
        search.search_input.setText("NEEDLE")
        assert search.matches == [0]
        assert not tab.table.isRowHidden(0) and tab.table.isRowHidden(1)
        assert tab.table.isRowHidden(2) and tab.table.isRowHidden(3)
        search.search_input.setText("Open")
        assert search.matches == [0, 3]
        assert "1 / 2" in search.count_label.text()
        search.next_button.click()
        assert tab.model.active_search_row == 3
        assert tab.model.data(tab.model.index(3, 0), Qt.ItemDataRole.BackgroundRole).color().name() == "#ffe08a"
        search.previous_button.click()
        assert tab.model.active_search_row == 0
        search.previous_button.click()
        assert tab.model.active_search_row == 3  # wrap
        search.clear_button.click()
        assert tab.model.active_search_row is None
        assert all(not tab.table.isRowHidden(r) for r in range(tab.model.rowCount()))
    assert not emitted
    assert p.model_dump() == data
    assert not other.is_populated and other.model.rowCount() == 0
    if edit_mode:
        tab.cancel_changes()
    editor.close()


@pytest.mark.parametrize("edit_mode", [False, True])
def test_search_filters_every_term_and_selects_only_visible_matches(qapp, edit_mode):
    p = make_project()
    first, second = p.units[0].components
    first.position = "C3"
    first.failure_mode_assignments[1].classification = "dangerous_failure"
    second.failure_mode_assignments[0].classification = "dangerous_failure"
    first.failure_mode_assignments[1].dc_test_ref = "Hidden isolated reference"
    window = MainWindow()
    window.current_project = p
    window.unit_editor_view.load_project(p)
    window.unit_editor_view.unit_tabs.setCurrentIndex(1)
    tab = window.unit_editor_view.unit_tabs.widget(1)
    other = window.unit_editor_view.unit_tabs.widget(2)
    if edit_mode:
        tab.enable_editing()
    search, table, model = tab.search_bar, tab.table, tab.model
    column_layout = [(table.isColumnHidden(c), table.columnWidth(c)) for c in range(model.columnCount())]
    row_order = [(e.component.id, e.fm_name) if not e.is_separator else None for e in model.rows]
    data = deepcopy(p.model_dump())
    window.has_unsaved_changes = False
    changes = []
    window.unit_editor_view.project_changed.connect(lambda: changes.append(True))

    def assert_filter(term, rows):
        search.search_input.setText(term)
        assert search.matches == rows
        assert [r for r in range(model.rowCount()) if not table.isRowHidden(r)] == rows
        assert table.isRowHidden(2)  # Separators never match.
        assert [(table.isColumnHidden(c), table.columnWidth(c)) for c in range(model.columnCount())] == column_layout

    with ExitStack() as guards:
        for method in ("calculate_project", "calculate_scope", "calculate_unit", "calculate_unit_metrics", "calculate_row_detailed"):
            guards.enter_context(patch.object(CalculationService, method, side_effect=AssertionError("Search recalculated")))
        for method in ("validate_row", "validate_project"):
            guards.enter_context(patch.object(ValidationService, method, side_effect=AssertionError("Search validated")))
        for target in (model, other.model):
            guards.enter_context(patch.object(target, "reload_data", side_effect=AssertionError("Search rebuilt table")))
        assert_filter("C3", [0, 1])
        assert_filter("Dangerous Failure", [1, 3])
        assert_filter("dAnGeRoUs fAiL", [1, 3])
        assert search.count_label.text() == "1 / 2 matches"
        for button, row in ((search.next_button, 3), (search.next_button, 1), (search.previous_button, 3)):
            button.click()
            assert table.currentIndex().row() == row
            assert [index.row() for index in table.selectionModel().selectedRows()] == [row]
            assert not table.isRowHidden(row)
            assert model.active_search_row == row
        assert_filter("hidden isolated", [1])
        assert table.isColumnHidden(16)
        # Component context remains available on a non-first failure-mode row.
        assert model.index(1, 0).data() == "C3"
        assert model.index(1, 3).data() == "Signal path"
        assert model.index(1, 5).data() == "Amplifier"
        assert [i.row() for i in table.selectionModel().selectedRows()] == [1]
        assert_filter("no-such-value", [])
        assert search.count_label.text() == "0 matches"
        assert not search.previous_button.isEnabled() and not search.next_button.isEnabled()
        assert not table.currentIndex().isValid() and not table.selectionModel().selectedRows()
        search.clear_button.click()
        assert all(not table.isRowHidden(r) for r in range(model.rowCount()))
        assert model.active_search_row is None
        assert [(e.component.id, e.fm_name) if not e.is_separator else None for e in model.rows] == row_order
    assert not changes and not window.undo_stack and not window.has_unsaved_changes
    assert p.model_dump() == data
    assert not other.is_populated and other.model.rowCount() == 0


@pytest.mark.parametrize("raw, normalized", [(None, None), ("", None), ("  ", None),
    ("yes", "Yes"), (" YES ", "Yes"), ("No", "No"), (True, "Yes"), (False, "No"),
    (1, "Yes"), (0, "No"), ("watchdog monitors CPU", None)])
def test_diagnostic_legacy_normalization(raw, normalized):
    a = FailureModeAssignment(failure_mode_name="Open", failure_rate_percentage=100, diagnostic_function=raw)
    assert a.diagnostic_function == normalized
    if raw == "watchdog monitors CPU":
        assert a.legacy_diagnostic_function == raw
    assert FailureModeAssignment.model_validate_json(a.model_dump_json()).model_dump() == a.model_dump()


def test_diagnostic_delegate_blank_and_independent_real_edits(qapp):
    p = make_project()
    editor = UnitEditorView()
    editor.load_project(p)
    editor.unit_tabs.setCurrentIndex(1)
    tab = editor.unit_tabs.widget(1)
    tab.enable_editing()
    model = tab.model
    idx = model.index(0, 10)
    delegate = tab.table.itemDelegateForColumn(10)
    widget = delegate.createEditor(None, QStyleOptionViewItem(), idx)
    assert isinstance(widget, QComboBox)
    delegate.setEditorData(widget, idx)
    assert [widget.itemText(i) for i in range(widget.count())] == ["", "Yes", "No"]
    assert widget.currentText() == ""
    events = []
    model.data_modified.connect(lambda: events.append(1))
    a = p.units[0].components[0].failure_mode_assignments[0]
    before = a.model_dump()
    assert not model.setData(idx, "")
    assert model.setData(idx, "Yes")
    assert not model.setData(idx, "Yes")
    assert not model.setData(idx, "text")
    assert a.model_dump() == dict(before, diagnostic_function="Yes")
    assert events == [1]
    assert "8.0000 FIT" in tab.diagnostic_summary.text()
    assert model.setData(idx, "No")
    assert "0.0000 FIT" in tab.diagnostic_summary.text()
    tab.confirm_changes()
    editor.close()


def test_diagnostic_persistence_warnings_and_scope(tmp_path, qapp):
    p = make_project()
    first = p.units[0].components[0]
    first.failure_mode_assignments[0].diagnostic_function = "Yes"
    first.failure_mode_assignments[0].dont_care = True
    first.failure_mode_assignments[1].diagnostic_function = "No"
    p.units[1].included_in_safety_function = False
    assert CalculationService.calculate_unit(p.units[0])["lambda_diag_gesamt"] == 8
    assert CalculationService.calculate_unit_metrics(p.units[1], p)["lambda_diag_gesamt"] == 10
    assert CalculationService.calculate_scope(p.units, p)["lambda_diag_gesamt"] == 18
    assert CalculationService.calculate_scope(p.units[:1], p)["lambda_diag_gesamt"] == 8
    assert CalculationService.calculate_scope([], p)["lambda_diag_gesamt"] == 0
    path = tmp_path / "diagnostic.json"
    ProjectService.save_project_atomically(p, str(path))
    loaded, _, _ = ProjectService.load_and_migrate_project(str(path))
    assert [a.diagnostic_function for a in loaded.units[0].components[0].failure_mode_assignments] == ["Yes", "No"]
    assert loaded.units[0].components[1].failure_mode_assignments[0].diagnostic_function is None
    alerts = [a for a in ValidationService.validate_project(loaded) if a["message"] == "Diagnostic Function not evaluated"]
    assert [(a["unit_id"], a["row_index"], a["severity"]) for a in alerts] == [("u1", 2, "Warning"), ("u1", 3, "Warning")]
    editor = UnitEditorView()
    editor.load_project(loaded)
    editor.focus_unit_row("u1", 2)
    tab = editor.unit_tabs.widget(1)
    assert tab.table.currentIndex().row() == 3  # separator skipped by warning indexing
    tab.search_bar.search_input.setText("R1")
    editor.focus_unit_row("u1", 3)
    assert tab.search_bar.search_input.text() == ""
    assert tab.table.currentIndex().row() == 4
    view = VerificationView()
    view.load_project(loaded)
    view._populate_secondary_table_from_scope(CalculationService.calculate_scope(p.units[:1], p))
    row = next(r for r in range(view.secondary_table.rowCount()) if "λDiagGesamt" in view.secondary_table.item(r, 0).text())
    assert view.secondary_table.item(row, 1).text() == "8.0000 FIT"


def test_diagnostic_undo_and_search_has_no_undo(qapp, fresh_isolated_db):
    window = MainWindow()
    p = make_project()
    window.current_project = p
    window.unit_editor_view.load_project(p)
    window.show_view("unit_editor")
    window.unit_editor_view.unit_tabs.setCurrentIndex(1)
    tab = window.unit_editor_view.unit_tabs.widget(1)
    window.has_unsaved_changes = False
    tab.search_bar.search_input.setText("Open")
    assert not window.undo_stack and not window.has_unsaved_changes
    tab.enable_editing()
    tab.confirm_changes()
    assert not window.undo_stack and not window.has_unsaved_changes
    tab.enable_editing()
    tab.model.setData(tab.model.index(0, 10), "Yes")
    tab.confirm_changes()
    assert len(window.undo_stack) == 1 and window.has_unsaved_changes
    window._on_undo()
    assert window.current_project.units[0].components[0].failure_mode_assignments[0].diagnostic_function is None
    assert CalculationService.calculate_unit(window.current_project.units[0])["lambda_diag_gesamt"] == 0


def test_save_audit_batches_changes_and_records_old_new(tmp_path):
    p = make_project()
    path = tmp_path / "audit.json"
    ProjectService.save_project_atomically(p, str(path), comment="Initial baseline")
    assert len(p.change_history) == 1
    p.units[0].name = "Renamed group"
    p.units[0].components[0].failure_mode_assignments[0].diagnostic_function = "Yes"
    p.units[0].components.pop()
    p.units.append(Unit(id="new", name="New group", description="Added"))
    ProjectService.save_project_atomically(p, str(path), comment="Engineering review")
    assert len(p.change_history) == 2
    entry = p.change_history[-1]
    assert entry["comment"] == "Engineering review"
    assert entry["user"] == "Engineer"
    assert len(entry["changes"]) == 4
    assert {c["action"] for c in entry["changes"]} == {"Added", "Changed", "Removed"}
    rename = next(c for c in entry["changes"] if c["field"].endswith(" / name"))
    assert rename["old_value"] == "Sensor group" and rename["new_value"] == "Renamed group"
    diag = next(c for c in entry["changes"] if c["field"].endswith(" / diagnostic function"))
    assert diag["old_value"] is None and diag["new_value"] == "Yes"
    assert json.loads(path.read_text(encoding="utf-8"))["change_history"] == p.change_history
    assert len(json.loads(path.with_suffix(".json.bak").read_text(encoding="utf-8"))["change_history"]) == 1
    assert "\n  " not in path.read_text(encoding="utf-8")


def test_no_audit_for_recalculation_ui_or_format_changes(tmp_path):
    p = make_project()
    path = tmp_path / "quiet.json"
    ProjectService.save_project_atomically(p, str(path))
    history = deepcopy(p.change_history)
    CalculationService.calculate_project(p)
    p.updated_at = datetime.now()
    p.units[0].updated_at = datetime.now()
    p.last_active_tab_id = "u2"
    ProjectService.save_project_atomically(p, str(path))
    assert p.change_history == history
    # Reopening an indented / differently ordered document is not a data change.
    path.write_text(json.dumps(p.model_dump(mode="json"), indent=4, sort_keys=True), encoding="utf-8")
    reopened, _, _ = ProjectService.load_and_migrate_project(str(path))
    ProjectService.save_project_atomically(reopened, str(path))
    assert reopened.change_history == history


def test_comment_edit_is_separate_and_persisted(tmp_path, qapp):
    p = make_project()
    path = tmp_path / "comment.json"
    ProjectService.save_project_atomically(p, str(path))
    facts = deepcopy(p.change_history[0])
    dialog = ChangeHistoryDialog(p)
    assert dialog.table.editTriggers() == QAbstractItemView.EditTrigger.NoEditTriggers
    assert dialog.details_table.editTriggers() == QAbstractItemView.EditTrigger.NoEditTriggers
    dialog.set_comment(0, "Reviewed\nClarified")
    assert p.change_history[0] == dict(facts, comment="Reviewed\nClarified")
    ProjectService.save_project_atomically(p, str(path))
    assert len(p.change_history) == 1
    loaded, _, _ = ProjectService.load_and_migrate_project(str(path))
    assert loaded.change_history == p.change_history


def test_failed_save_keeps_history_baseline_and_previous_file(tmp_path):
    p = make_project()
    path = tmp_path / "failure.json"
    ProjectService.save_project_atomically(p, str(path))
    valid_bytes = path.read_bytes()
    history, baseline = deepcopy(p.change_history), deepcopy(p._saved_state)
    p.reviewer = "New reviewer"
    with patch("fmeda_tool.services.project_service.os.replace", side_effect=OSError("disk failure")):
        with pytest.raises(OSError):
            ProjectService.save_project_atomically(p, str(path))
    assert path.read_bytes() == valid_bytes
    assert p.change_history == history and p._saved_state == baseline
    assert path.with_suffix(".json.bak").read_bytes() == valid_bytes
    assert not list(tmp_path.glob(".tmp*"))
    ProjectService.save_project_atomically(p, str(path))
    assert len(p.change_history) == 2
    assert p.change_history[-1]["changes"][0]["new_value"] == "New reviewer"


def test_save_as_baseline_and_main_window_path(tmp_path, qapp, fresh_isolated_db):
    p = make_project()
    original = tmp_path / "original.json"
    target = tmp_path / "copy.json"
    ProjectService.save_project_atomically(p, str(original))
    original_bytes = original.read_bytes()
    p.description = "Edited before Save As"
    window = MainWindow()
    window.current_project = p
    window.unit_editor_view.load_project(p)
    with patch("fmeda_tool.ui.main_window.QFileDialog.getSaveFileName", return_value=(str(target), "")), \
         patch.object(QInputDialog, "getMultiLineText", return_value=("Save As comment", True)), \
         patch.object(QMessageBox, "information"):
        window._on_save_as()
    assert original.read_bytes() == original_bytes
    assert p._saved_path == str(target.resolve())
    assert len(p.change_history) == 2
    with patch.object(QInputDialog, "getMultiLineText", side_effect=AssertionError("Unchanged save asked for comment")), \
         patch.object(QMessageBox, "information"):
        assert window._on_save_project_from_editor()
    assert len(p.change_history) == 2
    p.reviewer = "Reviewer"
    with patch.object(QInputDialog, "getMultiLineText", return_value=("", False)):
        assert window._on_save_project_from_editor() is False
    assert len(p.change_history) == 2


def test_large_save_has_one_entry_with_structured_details(tmp_path):
    p = make_project()
    path = tmp_path / "large.json"
    ProjectService.save_project_atomically(p, str(path))
    p.custom_fields = {f"Engineering {i}": str(i) for i in range(250)}
    ProjectService.save_project_atomically(p, str(path))
    assert len(p.change_history) == 2
    assert len(p.change_history[-1]["changes"]) == 250
    assert len(p.change_history[-1]["details"]) < 400


def test_diagnostic_and_history_exports(tmp_path, qapp):
    p = make_project()
    p.units[0].components[0].failure_mode_assignments[0].diagnostic_function = "Yes"
    ProjectService.save_project_atomically(p, str(tmp_path / "p.json"), comment="Exported comment")
    xlsx = tmp_path / "export.xlsx"
    assert ExportService.export_to_excel(p, str(xlsx), include_history=True)
    book = openpyxl.load_workbook(xlsx)
    overview = book["Overview"]
    label_cells = [cell for row in overview for cell in row if cell.value == "λDiagGesamt (FIT)"]
    assert len(label_cells) == 2
    header = next(cell for cell in label_cells if cell.column > 3)
    assert overview.cell(header.row + 1, header.column).value == "8.0000"
    assert overview.cell(header.row + 2, header.column).value == "10.0000"
    assert book["Change History"].cell(2, 8).value == "Exported comment"
    assert book["Change Details"].max_row == 2
    assert all(c.width <= 70 for c in book["Change History"].column_dimensions.values())
    assert book["Change Details"].cell(2, 8).alignment.wrap_text
    book.close()
    assert ExportService.export_to_excel(p, str(xlsx), include_history=False)
    book = openpyxl.load_workbook(xlsx)
    assert "Change History" not in book.sheetnames and "Change Details" not in book.sheetnames
    book.close()
    captured = []
    original = QTextDocument.setHtml
    def capture(doc, html):
        original(doc, html)
        captured.append(doc.toPlainText())
    with patch.object(QTextDocument, "setHtml", capture):
        assert ExportService.export_to_pdf(p, str(tmp_path / "export.pdf"))
    assert "λDiagGesamt" in captured[0]
    assert "18.0000 FIT" in captured[0] and "8.0000 FIT" in captured[0] and "10.0000 FIT" in captured[0]


def test_search_edit_invalidates_index_and_selection_stays_visible(qapp):
    p = make_project()
    editor = UnitEditorView()
    editor.load_project(p)
    editor.unit_tabs.setCurrentIndex(1)
    tab = editor.unit_tabs.widget(1)
    tab.enable_editing()
    tab.search_bar.search_input.setText("R1")
    tab.table.select_all_components()
    assert tab.table.selected_component_ids == {"c0"}
    tab.search_bar.search_input.setText("Row justification")
    assert tab.search_bar.matches == [0, 1, 3, 4]
    tab.model.setData(tab.model.index(0, 18), "Replaced note")
    assert tab.search_bar.matches == [1, 3, 4]
    tab.search_bar.search_input.setText("replaced")
    assert tab.search_bar.matches == [0]
    assert tab.table.currentIndex().row() == 0
    tab.cancel_changes()
    editor.close()


@pytest.mark.parametrize("field, new_value", [
    ("classification", "dangerous_failure"), ("dont_care", True), ("proof_test_a", 70),
    ("proof_test_b", 80), ("proof_test_c", 90), ("deviation_id", "dev-new"),
    ("diagnostic_measure_id", "dm-new"), ("mitigation_id", "mit-new"),
    ("failure_rate_percentage", 45), ("detection_percentage", 95), ("notes", "Changed"),
])
def test_engineering_assignment_changes_audited(tmp_path, field, new_value):
    p = make_project()
    path = tmp_path / "engineering.json"
    ProjectService.save_project_atomically(p, str(path))
    a = p.units[0].components[0].failure_mode_assignments[0]
    old = getattr(a, field)
    setattr(a, field, new_value)
    ProjectService.save_project_atomically(p, str(path))
    assert len(p.change_history) == 2
    change, = p.change_history[-1]["changes"]
    assert change["old_value"] == old and change["new_value"] == new_value
    assert change["field"].endswith(field.replace("_", " "))


def test_older_project_blank_assignment_and_lazy_loading_do_not_create_audit(tmp_path, qapp):
    p = make_project()
    p.units[0].components[0].failure_mode_assignments = []
    data = p.model_dump(mode="json")
    for unit in data["units"]:
        for comp in unit["components"]:
            for assignment in comp["failure_mode_assignments"]:
                assignment.pop("diagnostic_function", None)
                assignment.pop("legacy_diagnostic_function", None)
    path = tmp_path / "old.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    loaded, _, _ = ProjectService.load_and_migrate_project(str(path))
    assert loaded.units[1].components[0].failure_mode_assignments[0].diagnostic_function is None
    alerts = ValidationService.validate_project(loaded)
    assert sum(a["message"] == "Diagnostic Function not evaluated" for a in alerts) == 5
    FmedaTableModel(loaded.units[0], loaded)
    CalculationService.calculate_project(loaded)
    ProjectService.save_project_atomically(loaded, str(path))
    assert loaded.change_history == []


def test_overview_edits_are_batched_only_at_save(tmp_path, qapp, fresh_isolated_db):
    p = make_project()
    path = tmp_path / "overview.json"
    ProjectService.save_project_atomically(p, str(path))
    editor = UnitEditorView()
    editor.load_project(p)
    overview = editor.unit_tabs.widget(0)
    overview.reviewer_input.setText("Reviewer A")
    overview._on_reviewer_changed()
    overview.reviewer_input.setText("Reviewer B")
    overview._on_reviewer_changed()
    assert len(p.change_history) == 1
    ProjectService.save_project_atomically(p, str(path))
    assert len(p.change_history) == 2
    change, = p.change_history[-1]["changes"]
    assert change["old_value"] is None and change["new_value"] == "Reviewer B"


def test_history_comment_marks_dirty_without_changing_audit_facts(tmp_path, qapp, fresh_isolated_db):
    p = make_project()
    path = tmp_path / "comments.json"
    ProjectService.save_project_atomically(p, str(path))
    facts = deepcopy(p.change_history[0])
    window = MainWindow()
    window.current_project = p
    window.unit_editor_view.load_project(p)
    overview = window.unit_editor_view.unit_tabs.widget(0)
    window.has_unsaved_changes = False
    with patch.object(ChangeHistoryDialog, "exec", lambda dialog: dialog.set_comment(0, "Reviewed change")):
        overview._on_view_history()
    assert window.has_unsaved_changes
    assert {k: v for k, v in p.change_history[0].items() if k != "comment"} == {k: v for k, v in facts.items() if k != "comment"}
    ProjectService.save_project_atomically(p, str(path))
    assert len(p.change_history) == 1
    assert json.loads(path.read_text(encoding="utf-8"))["change_history"][0]["comment"] == "Reviewed change"
    window.has_unsaved_changes = False
    with patch.object(ChangeHistoryDialog, "exec", lambda dialog: dialog.set_comment(0, "Reviewed change")):
        overview._on_view_history()
    assert not window.has_unsaved_changes
