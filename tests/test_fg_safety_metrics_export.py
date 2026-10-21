import unittest
import os
import tempfile
from pathlib import Path
from typing import List, Dict, Any

from PyQt6.QtWidgets import QApplication, QLabel, QFrame, QPushButton
from openpyxl import load_workbook

from fmeda_tool.models import (
    Project, Unit, Component, FailureModeAssignment,
    ComponentDB, SafetyStandard, ProjectStatus, SafetyContext
)
from fmeda_tool.services import CalculationService, ExportService
from fmeda_tool.ui.dialogs.component_mapping_dialog import ComponentMappingDialog

app = QApplication.instance()
if app is None:
    app = QApplication([])


class TestFunctionalGroupSafetyMetricsExport(unittest.TestCase):
    """
    Test suite for:
    1. BOM mapping workflow instruction layout and section headings.
    2. Independent functional group safety calculations.
    3. Comprehensive Excel export of individual functional group safety metrics.
    """

    def setUp(self):
        # Unit 1: SFTY-ADC-CPU (Safety Function = True)
        self.comp_adc = Component(
            id="comp_adc_01",
            name="Microcontroller ADC",
            position="IC1",
            type="Microcontroller",
            failure_rate=100.0,
            failure_modes={"Hardware Fault": 60.0, "Drift": 40.0},
            failure_mode_assignments=[
                FailureModeAssignment(
                    failure_mode_name="Hardware Fault",
                    failure_rate_percentage=60.0,
                    classification="dangerous",
                    dangerous_failure_percentage=100.0,
                    detection_percentage=90.0,  # 60 * 1.0 * 0.9 = 54 λdd, 6 λdu
                    proof_test_a=95.0,
                    proof_test_b=80.0,
                    proof_test_c=50.0,
                    dont_care=False
                ),
                FailureModeAssignment(
                    failure_mode_name="Drift",
                    failure_rate_percentage=40.0,
                    classification="safe",
                    dangerous_failure_percentage=0.0,  # 40 * 1.0 = 40 λs
                    detection_percentage=50.0,  # 20 λsd, 20 λsu
                    proof_test_a=0.0,
                    proof_test_b=0.0,
                    proof_test_c=0.0,
                    dont_care=False
                )
            ]
        )
        self.unit_adc = Unit(
            id="unit_adc",
            name="SFTY-ADC-CPU",
            description="ADC & CPU Channel",
            included_in_safety_function=True,
            components=[self.comp_adc]
        )

        # Unit 2: SFTY-AO (Safety Function = True, with a Don't Care assignment)
        self.comp_ao = Component(
            id="comp_ao_01",
            name="Analog Output Driver",
            position="U2",
            type="OpAmp",
            failure_rate=50.0,
            failure_modes={"Output Stuck High": 50.0, "Output Stuck Low": 50.0},
            failure_mode_assignments=[
                FailureModeAssignment(
                    failure_mode_name="Output Stuck High",
                    failure_rate_percentage=50.0,
                    classification="dangerous",
                    dangerous_failure_percentage=100.0,
                    detection_percentage=80.0,  # 25 * 0.8 = 20 λdd, 5 λdu
                    proof_test_a=99.0,
                    proof_test_b=90.0,
                    proof_test_c=0.0,
                    dont_care=False
                ),
                FailureModeAssignment(
                    failure_mode_name="Output Stuck Low",
                    failure_rate_percentage=50.0,
                    classification="dangerous",
                    dangerous_failure_percentage=100.0,
                    detection_percentage=80.0,
                    proof_test_a=99.0,
                    proof_test_b=90.0,
                    proof_test_c=0.0,
                    dont_care=True  # DON'T CARE row: excluded from Safety Channel!
                )
            ]
        )
        self.unit_ao = Unit(
            id="unit_ao",
            name="SFTY-AO",
            description="Analog Output Channel",
            included_in_safety_function=True,
            components=[self.comp_ao]
        )

        # Unit 3: AUX-POWER (Non-Safety Unit, included_in_safety_function = False)
        self.comp_aux = Component(
            id="comp_aux_01",
            name="Auxiliary Regulator",
            position="VR1",
            type="Voltage Regulator",
            failure_rate=30.0,
            failure_modes={"Output Loss": 100.0},
            failure_mode_assignments=[
                FailureModeAssignment(
                    failure_mode_name="Output Loss",
                    failure_rate_percentage=100.0,
                    classification="safe",
                    dangerous_failure_percentage=0.0,
                    detection_percentage=0.0,
                    dont_care=False
                )
            ]
        )
        self.unit_aux = Unit(
            id="unit_aux",
            name="AUX-POWER",
            description="Auxiliary Power Supply",
            included_in_safety_function=False,
            components=[self.comp_aux]
        )

        self.project = Project(
            id="proj_export_test",
            name="Safety Metrics Export Test Project",
            description="Project verifying independent group calculations and export",
            status=ProjectStatus.DRAFT,
            safety_standard=SafetyStandard.IEC_61508,
            target_sil="SIL 2",
            test_interval=8760.0,
            diagnostic_test_interval=8.0,
            selected_profile="Profile 1",
            safety_context=SafetyContext(
                safety_function_name="Overpressure Protection",
                safety_architecture="1oo1"
            ),
            units=[self.unit_adc, self.unit_ao, self.unit_aux]
        )

    # ========================================================
    # 1. Workflow Readability & UI Layout Tests
    # ========================================================

    def test_workflow_cards_layout_and_labels(self):
        """Test that the integrated workspace renders with clear header, sections, and workflow help button."""
        dialog = ComponentMappingDialog(self.unit_adc, project_profile="Profile 1")
        
        # Verify dialog window title and profile
        self.assertIn("SFTY-ADC-CPU", dialog.windowTitle())
        self.assertIn("Profile 1", dialog.windowTitle())
        
        # Check section labels
        labels = dialog.findChildren(QLabel)
        label_texts = [lbl.text() for lbl in labels]
        
        # Section titles
        self.assertTrue(any("1. Component Library" in t for t in label_texts))
        self.assertTrue(any("2. BOM Positions" in t for t in label_texts))
        self.assertTrue(any("3. Mapping Summary" in t for t in label_texts))
        
        # Workflow help button exists and has tooltip
        buttons = dialog.findChildren(QPushButton)
        help_buttons = [b for b in buttons if "Workflow Help" in b.text()]
        self.assertEqual(len(help_buttons), 1)
        self.assertIn("Workflow Help", help_buttons[0].toolTip())

    def test_workflow_dialog_resizing(self):
        """Test that dialog can be resized without clipping or crashing."""
        dialog = ComponentMappingDialog(self.unit_adc, project_profile="Profile 1")
        dialog.resize(1000, 600)
        dialog.resize(1400, 900)
        dialog.resize(1180, 720)
        self.assertEqual(dialog.width(), 1180)
        self.assertEqual(dialog.height(), 720)

    # ========================================================
    # 2. Independent Functional Group Safety Calculations Tests
    # ========================================================

    def test_independent_group_calculations_sfty_adc_cpu(self):
        """Test that SFTY-ADC-CPU is calculated purely from its own components and assignments."""
        m = CalculationService.calculate_unit_metrics(self.unit_adc, self.project)
        
        self.assertEqual(m["unit_name"], "SFTY-ADC-CPU")
        self.assertTrue(m["included_in_safety_function"])
        
        # Gesamtgerät: 100 FIT total (40 Safe: 20 sd, 20 su; 60 Dangerous: 54 dd, 6 du)
        self.assertAlmostEqual(m["lambda_total_gesamtgerat"], 100.0, places=4)
        self.assertAlmostEqual(m["lambda_safe_gesamtgerat"], 40.0, places=4)
        self.assertAlmostEqual(m["lambda_dangerous_gesamtgerat"], 60.0, places=4)
        self.assertAlmostEqual(m["lambda_sd_gesamtgerat"], 20.0, places=4)
        self.assertAlmostEqual(m["lambda_su_gesamtgerat"], 20.0, places=4)
        self.assertAlmostEqual(m["lambda_dd_gesamtgerat"], 54.0, places=4)
        self.assertAlmostEqual(m["lambda_du_gesamtgerat"], 6.0, places=4)
        
        # SFF Gesamtgerät = (40 + 54) / 100 * 100 = 94.0%
        self.assertAlmostEqual(m["sff_gesamtgerat"], 94.0, places=2)
        
        # DC Gesamtgerät = 54 / 60 * 100 = 90.0%
        self.assertAlmostEqual(m["dc_gesamtgerat"], 90.0, places=2)
        
        # Sicherheitskanal (no don't care in adc): same rates
        self.assertAlmostEqual(m["lambda_total_sicherheitskanal"], 100.0, places=4)
        self.assertAlmostEqual(m["lambda_dd_sicherheitskanal"], 54.0, places=4)
        self.assertAlmostEqual(m["lambda_du_sicherheitskanal"], 6.0, places=4)
        self.assertAlmostEqual(m["sff_sicherheitskanal"], 94.0, places=2)
        self.assertAlmostEqual(m["dc_sicherheitskanal"], 90.0, places=2)
        
        # MTTFd Sicherheitskanal = 10^9 / (60 * 8760) = 1902.58 years
        expected_mttfd = 10**9 / (60.0 * 8760.0)
        self.assertAlmostEqual(m["mttfd_sicherheitskanal"], expected_mttfd, places=2)
        
        # Achieved SIL (SFF 94% -> SIL 2)
        self.assertEqual(m["achieved_sil"], "SIL 2")
        
        # Proof Test A Coverage over lambda_du: (6 * 0.95) / 6 = 95.0%
        self.assertAlmostEqual(m["proof_test_a_coverage"], 95.0, places=2)
        self.assertAlmostEqual(m["proof_test_b_coverage"], 80.0, places=2)
        self.assertAlmostEqual(m["proof_test_c_coverage"], 50.0, places=2)

    def test_independent_group_calculations_sfty_ao_dont_care(self):
        """Test that SFTY-AO DON'T CARE exclusion affects ONLY the safety channel."""
        m = CalculationService.calculate_unit_metrics(self.unit_ao, self.project)
        
        self.assertEqual(m["unit_name"], "SFTY-AO")
        self.assertTrue(m["included_in_safety_function"])
        
        # Gesamtgerät includes BOTH failure modes (50 FIT total, 40 dd, 10 du)
        self.assertAlmostEqual(m["lambda_total_gesamtgerat"], 50.0, places=4)
        self.assertAlmostEqual(m["lambda_dd_gesamtgerat"], 40.0, places=4)
        self.assertAlmostEqual(m["lambda_du_gesamtgerat"], 10.0, places=4)
        self.assertAlmostEqual(m["sff_gesamtgerat"], 80.0, places=2)
        
        # Sicherheitskanal EXCLUDES the DON'T CARE assignment (25 FIT total, 20 dd, 5 du)
        self.assertAlmostEqual(m["lambda_total_sicherheitskanal"], 25.0, places=4)
        self.assertAlmostEqual(m["lambda_dd_sicherheitskanal"], 20.0, places=4)
        self.assertAlmostEqual(m["lambda_du_sicherheitskanal"], 5.0, places=4)
        self.assertAlmostEqual(m["sff_sicherheitskanal"], 80.0, places=2)
        self.assertAlmostEqual(m["dc_sicherheitskanal"], 80.0, places=2)
        
        # Achieved SIL (SFF 80% -> SIL 1)
        self.assertEqual(m["achieved_sil"], "SIL 1")
        
        # Proof test coverages
        self.assertAlmostEqual(m["proof_test_a_coverage"], 99.0, places=2)
        self.assertAlmostEqual(m["proof_test_b_coverage"], 90.0, places=2)
        self.assertAlmostEqual(m["proof_test_c_coverage"], 0.0, places=2)

    def test_non_safety_unit_metrics_show_none_or_na(self):
        """Test that non-safety functional groups have None / N/A for safety channel metrics and SIL."""
        m = CalculationService.calculate_unit_metrics(self.unit_aux, self.project)
        
        self.assertEqual(m["unit_name"], "AUX-POWER")
        self.assertFalse(m["included_in_safety_function"])
        
        # Gesamtgerät exists
        self.assertAlmostEqual(m["lambda_total_gesamtgerat"], 30.0, places=4)
        self.assertAlmostEqual(m["lambda_safe_gesamtgerat"], 30.0, places=4)
        self.assertAlmostEqual(m["sff_gesamtgerat"], 100.0, places=2)
        
        # Sicherheitskanal is not applicable
        self.assertIsNone(m["sff_sicherheitskanal"])
        self.assertIsNone(m["dc_sicherheitskanal"])
        self.assertIsNone(m["mttfd_sicherheitskanal"])
        self.assertIsNone(m["pfd_avg"])
        self.assertIsNone(m["pfd_max"])
        self.assertEqual(m["achieved_sil"], "N/A")

    # ========================================================
    # 3. Excel Export Structure & Content Tests
    # ========================================================

    def test_excel_export_contains_independent_fg_metrics_table(self):
        """Test that the exported Excel workbook contains the 24-column Calculated Safety Metrics by Functional Group table."""
        with tempfile.TemporaryDirectory() as tmp_dir:
            file_path = os.path.join(tmp_dir, "test_fmeda_export.xlsx")
            success = ExportService.export_to_excel(self.project, file_path)
            self.assertTrue(success)
            self.assertTrue(os.path.exists(file_path))
            
            wb = load_workbook(file_path, data_only=True)
            self.assertIn("Overview", wb.sheetnames)
            ws = wb["Overview"]
            
            # Find the section title "Calculated Safety Metrics by Functional Group"
            found_title = False
            title_row = -1
            for row in range(1, 100):
                val = ws.cell(row=row, column=1).value
                if val and "Calculated Safety Metrics by Functional Group" in str(val):
                    found_title = True
                    title_row = row
                    break
                    
            self.assertTrue(found_title, "Calculated Safety Metrics by Functional Group section title not found")
            
            # Check headers row (title_row + 2)
            header_row = title_row + 2
            expected_headers = ["Functional Group"] + [label for _, label, _ in CalculationService.GROUP_METRICS]
            self.assertEqual([ws.cell(header_row, c).value for c in range(1, len(expected_headers) + 1)], expected_headers)
            for offset, unit in enumerate(self.project.units, 1):
                result = CalculationService.calculate_group_result(unit, self.project)
                self.assertEqual(ws.cell(header_row + offset, 1).value, unit.name)
                for col, (key, _, style) in enumerate(CalculationService.GROUP_METRICS, 2):
                    expected = result[key]
                    if expected is None:
                        expected = "N/A"
                    elif style == "pct":
                        expected /= 100
                    actual = ws.cell(header_row + offset, col).value
                    if isinstance(expected, float):
                        self.assertAlmostEqual(actual, expected)
                    else:
                        self.assertEqual(actual, expected)
            self.assertFalse(any(c.value == "Calculated Safety Metrics Summary" for row in ws for c in row))
            self.assertEqual(ws.max_row, header_row + len(self.project.units))
            wb.close()


if __name__ == "__main__":
    unittest.main()
